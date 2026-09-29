"""Full case review with inline highlighting of changed text."""
import difflib
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

source = json.loads(Path('/Users/stanislav/Downloads/Редактор публикаций.json').read_text(encoding='utf-8'))
output = json.loads(Path('artifacts/Редактор публикаций — актуализировано.json').read_text(encoding='utf-8'))
old_cases = {case['id']: case for case in source['suites'][0]['test_cases']}
cases = output['suites'][0]['test_cases']
target = Path('artifacts/Редактор публикаций — обзор изменений.docx')

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.left_margin = section.right_margin = Inches(.82)
section.top_margin = section.bottom_margin = Inches(.68)
for name in ('Normal', 'Title', 'Heading 2'):
    doc.styles[name].font.name = 'Arial'
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(9.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(4)
doc.styles['Title'].font.size = Pt(18)
title_ppr = doc.styles['Title']._element.get_or_add_pPr()
for border in title_ppr.findall(qn('w:pBdr')):
    title_ppr.remove(border)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(13)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 2'].paragraph_format.keep_with_next = True


def shade_run(run, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    run._r.get_or_add_rPr().append(shd)


def add_diff(paragraph, before, after, new=False):
    before = '—' if before is None or before == '' else str(before)
    after = '—' if after is None or after == '' else str(after)
    if new:
        run = paragraph.add_run(after)
        run.font.color.rgb = RGBColor(60, 108, 77)
        shade_run(run, 'EAF5EC')
        return
    old_tokens = re.findall(r'\S+\s*|\s+', before)
    new_tokens = re.findall(r'\S+\s*|\s+', after)
    for op, _, _, start, end in difflib.SequenceMatcher(None, old_tokens, new_tokens, autojunk=False).get_opcodes():
        if op == 'delete':
            continue
        run = paragraph.add_run(''.join(new_tokens[start:end]))
        if op != 'equal':
            run.font.color.rgb = RGBColor(57, 94, 119)
            shade_run(run, 'EAF3F8')


def field(label, old, current, new=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run(label + '  ').bold = True
    add_diff(paragraph, old, current, new)


doc.add_paragraph('Кейсы редактора публикаций', 'Title')
existing_count = sum('id' in case for case in cases)
new_count = len(cases) - existing_count
doc.add_paragraph(f'Полная версия набора для чтения перед импортом. Показаны все {len(cases)} кейсов: {existing_count} исходных и {new_count} новых. Кейс 7684 целиком удалён: произвольное имя для публикации больше не задаётся. Светло-голубым выделен только добавленный или изменённый текст, зелёным — новые кейсы. Неподсвеченный текст сохранён из исходного экспорта; удалённые фрагменты в итоговой версии не показываются.')
doc.add_paragraph('Для фото и видео в предусловии задана настройка админки 100 МБ (диапазон настройки 1–2048 МБ). Для аудио и документов предел фиксированный — 100 МБ. В 7666 проверяется именно возможность выбрать 11 файлов одного типа одной операцией, без предположения о полном отсутствии лимита. Новым кейсам не присвоены ID и ключи: их должен выдать проект при импорте. Новые проверки панели инструментов имеют статус draft до фактического прогона во всех трёх представлениях.')
doc.add_paragraph('Ссылки покрыты отдельными сценариями: выделенный текст, URL без выделения, фрагмент заголовка, поиск публикации и выбор со следующей страницы результатов. Для них предусмотрены проверка в предпросмотре и переход из опубликованной записи. Текстовые ссылки в заголовке доступны даже тогда, когда инструменты форматирования заголовка отключены.')

for case in cases:
    old = old_cases.get(case.get('id'))
    is_new = old is None
    heading = doc.add_paragraph(style='Heading 2')
    if is_new:
        run = heading.add_run('НОВЫЙ КЕЙС  ')
        run.font.color.rgb = RGBColor(60, 108, 77)
        shade_run(run, 'EAF5EC')
    else:
        heading.add_run(f"{case['id']}  ")
    add_diff(heading, old.get('title') if old else None, case.get('title'), is_new)
    meta = doc.add_paragraph()
    meta.add_run('Статус  ').bold = True
    add_diff(meta, old.get('status') if old else None, case.get('status'), is_new)
    meta.add_run('    Приоритет  ').bold = True
    add_diff(meta, old.get('priority') if old else None, case.get('priority'), is_new)
    field('Описание', old.get('description') if old else None, case.get('description'), is_new)
    field('Предусловия', old.get('preconditions') if old else None, case.get('preconditions'), is_new)
    old_steps = old.get('steps', []) if old else []
    new_steps = case.get('steps', [])
    label = doc.add_paragraph()
    label.add_run('Шаги').bold = True
    if not new_steps:
        label.add_run('  —')
    old_signatures = [(s.get('action'), s.get('expected_result')) for s in old_steps]
    new_signatures = [(s.get('action'), s.get('expected_result')) for s in new_steps]
    aligned = {}
    if not is_new:
        matcher = difflib.SequenceMatcher(None, old_signatures, new_signatures, autojunk=False)
        for op, old_start, old_end, new_start, new_end in matcher.get_opcodes():
            if op == 'equal':
                for offset in range(new_end - new_start):
                    aligned[new_start + offset] = old_steps[old_start + offset]
            elif op == 'replace':
                for offset in range(min(old_end - old_start, new_end - new_start)):
                    aligned[new_start + offset] = old_steps[old_start + offset]
    for index, step in enumerate(new_steps):
        old_step = aligned.get(index, {})
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.16)
        p.paragraph_format.keep_together = True
        p.add_run(f"{index + 1}.  ").bold = True
        add_diff(p, old_step.get('action'), step.get('action'), is_new)
        p.add_run('\nОжидается:  ')
        add_diff(p, old_step.get('expected_result'), step.get('expected_result'), is_new)
    field('Постусловия', old.get('postconditions') if old else None, case.get('postconditions'), is_new)

footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Редактор публикаций  •  полная версия с изменениями')
doc.save(target)
print(f'DOCX created: {len(cases)} cases')
