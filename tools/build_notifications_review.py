"""Readable full-case Word review with changes against the supplied export."""
import difflib
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

source = json.loads(Path('/Users/stanislav/Downloads/Уведомления.json').read_text(encoding='utf-8'))
output = json.loads(Path('artifacts/Уведомления — актуализировано.json').read_text(encoding='utf-8'))
old_cases = {case['id']: case for case in source['suites'][0]['test_cases']}
cases = output['suites'][0]['test_cases']
target = Path('artifacts/Уведомления — обзор изменений.docx')

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.left_margin = section.right_margin = Inches(.82)
section.top_margin = section.bottom_margin = Inches(.58)
for name in ('Normal', 'Title', 'Heading 2'):
    doc.styles[name].font.name = 'Arial'
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(9.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(3)
doc.styles['Title'].font.size = Pt(18)
title_ppr = doc.styles['Title']._element.get_or_add_pPr()
for border in title_ppr.findall(qn('w:pBdr')):
    title_ppr.remove(border)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(13)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 2'].paragraph_format.keep_with_next = True


def shade_run(run, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    run._r.get_or_add_rPr().append(shd)


def add_diff(paragraph, before, after):
    before = '—' if before is None or before == '' else str(before)
    after = '—' if after is None or after == '' else str(after)
    old_tokens = re.findall(r'\S+\s*|\s+', before)
    new_tokens = re.findall(r'\S+\s*|\s+', after)
    for op, _, _, start, end in difflib.SequenceMatcher(None, old_tokens, new_tokens, autojunk=False).get_opcodes():
        if op == 'delete':
            continue
        run = paragraph.add_run(''.join(new_tokens[start:end]))
        if op != 'equal':
            run.font.color.rgb = RGBColor(57, 94, 119)
            shade_run(run, 'EAF3F8')


def field(label, old, current):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run(label + '  ').bold = True
    add_diff(paragraph, old, current)


doc.add_paragraph('Уведомления полная версия кейсов', 'Title')
doc.add_paragraph(
    'Все 10 оставшихся кейсов показаны целиком. Кейс 7464 о значке уведомлений и обновлении страницы исключён как неприменимый. Светло голубым выделен добавленный или изменённый текст. '
    'Неподсвеченные фрагменты сохранены из исходного экспорта; удалённый текст в итоговую версию не включён. '
    'Кейсы 7479 и 7472 были пустыми и теперь содержат полные сценарии.'
)
doc.add_paragraph(
    'В наборе проверяются только письма: их содержание и переходы по ссылкам на группу, пользователя и публикацию.'
)

for case in cases:
    old = old_cases[case['id']]
    heading = doc.add_paragraph(style='Heading 2')
    heading.add_run(f"{case['id']}  ")
    add_diff(heading, old.get('title'), case.get('title'))
    meta = doc.add_paragraph()
    meta.paragraph_format.keep_with_next = True
    meta.add_run('Статус  ').bold = True
    add_diff(meta, old.get('status'), case.get('status'))
    meta.add_run('    Приоритет  ').bold = True
    add_diff(meta, old.get('priority'), case.get('priority'))
    field('Описание', old.get('description'), case.get('description'))
    field('Предусловия', old.get('preconditions'), case.get('preconditions'))
    old_steps = old.get('steps', [])
    new_steps = case.get('steps', [])
    label = doc.add_paragraph()
    label.add_run('Шаги').bold = True
    if not new_steps:
        label.add_run('  —')
    for index, step in enumerate(new_steps):
        old_step = old_steps[index] if index < len(old_steps) else {}
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(.16)
        paragraph.paragraph_format.keep_together = True
        paragraph.add_run(f"{index + 1}.  ").bold = True
        add_diff(paragraph, old_step.get('action'), step.get('action'))
        paragraph.add_run('\nОжидается:  ')
        add_diff(paragraph, old_step.get('expected_result'), step.get('expected_result'))
    field('Постусловия', old.get('postconditions'), case.get('postconditions'))

footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Уведомления  •  полная версия с изменениями')
doc.save(target)
print(f'DOCX created: {len(cases)} cases')
