import copy
import difflib
import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path('/Users/stanislav/Downloads/Все кейсы.json')
OUT = Path(__file__).parent
original = json.loads(SOURCE.read_text(encoding='utf-8'))
updated = copy.deepcopy(original)
changed = {}
added = []


def suites(data):
    for suite in data['suites']:
        yield suite
        for nested in suite.get('suites', []):
            yield nested


def find_case(case_id):
    for suite in suites(updated):
        for test_case in suite.get('test_cases', []):
            if test_case.get('id') == case_id:
                changed[case_id] = copy.deepcopy(test_case)
                return test_case
    raise ValueError(case_id)


def steps(items):
    return [dict(position=i, action=action, expected_result=expected)
            for i, (action, expected) in enumerate(items, 1)]


c = find_case(7429)
c['title'] = 'Переключение между десктопным и мобильным видом'
c['description'] = 'Проверка перестроения интерфейса при переходе от широкой области просмотра к ширине смартфона. Точное значение breakpoint не фиксируется этим кейсом.'
c['preconditions'] = 'Пользователь авторизован; открыта лента соцсети в браузере с возможностью менять ширину области просмотра.'
c['steps'] = steps([
    ('Установить ширину области просмотра 1366 px и открыть ленту.', 'Отображается десктопная компоновка: навигация и карточки не перекрывают друг друга.'),
    ('Установить ширину 390 px и обновить страницу.', 'Интерфейс перестраивается под смартфон; навигация, карточки и основные действия доступны без горизонтальной прокрутки страницы.'),
    ('Плавно изменить ширину между 390 и 1366 px.', 'При переходе между компоновками контент не теряется, элементы не перекрываются; после возврата к 1366 px снова показан десктопный вид.'),
])

c = find_case(7542)
c['steps'][7]['expected_result'] = 'Выбранный дивизион отображается в поле; после заполнения обязательных полей доступна кнопка «Создать».'
c['steps'][9]['action'] = 'Нажать кнопку «Создать»'
c['steps'][9]['expected_result'] = 'Появляется сообщение «Группа успешно создана»; новая группа доступна в разделе «Группы» на вкладке «Вы автор», введённые данные сохранены.'

c = find_case(7603)
c['title'] = 'Восстановление позиции в списке групп после возврата'
c['description'] = 'После открытия группы и возврата в список сохраняется позиция прокрутки списка групп.'
c['steps'][0]['expected_result'] = 'Видны группы ниже начальных карточек списка.'
c['steps'][1]['expected_result'] = 'Открывается выбранная группа.'

c = find_case(10339)
c['title'] = 'Выбор места при создании публикации'
c['description'] = 'Проверка выбора группы и личной страницы при создании двух отдельных публикаций.'
c['preconditions'] = 'Пользователь авторизован; создана тестовая группа, в которой ему разрешено публиковать; открыт редактор новой публикации.'
c['steps'] = steps([
    ('В селекторе места публикации выбрать тестовую группу.', 'В шапке редактора указана выбранная группа.'),
    ('Ввести заголовок и нажать «Сейчас».', 'Публикация создана в выбранной группе.'),
    ('Открыть группу в соцсети.', 'Новая публикация отображается в группе.'),
    ('Открыть новый редактор через «Написать» и выбрать в селекторе «Моя лента».', 'В шапке редактора указана личная страница пользователя.'),
    ('Ввести другой заголовок и нажать «Сейчас».', 'Вторая публикация создана на личной странице.'),
    ('Открыть «Моя страница».', 'Вторая публикация отображается на личной странице и не является записью тестовой группы.'),
])

c = find_case(10370)
c['description'] = 'Проверка, что после подтверждённого закрытия без сохранения несохранённый контент не восстанавливается. Поведение при обновлении страницы проверяется отдельно в этом же кейсе.'
c['steps'][1]['action'] = 'Закрыть редактор крестиком и в запросе «Сохранить изменения?» выбрать «Закрыть и удалить изменения».'
c['steps'][1]['expected_result'] = 'Редактор закрывается без сохранения черновика.'
c['steps'][2]['expected_result'] = 'Несохранённые заголовок и текст отсутствуют в новом редакторе.'
c['steps'][3]['expected_result'] = 'После обновления страницы и повторного открытия редактора несохранённые заголовок и текст отсутствуют.'

c = find_case(10372)
c['steps'] = c['steps'][:3]


def add(suite_title, title, description, preconditions, pairs, status='actual'):
    suite = next(s for s in suites(updated) if s['title'] == suite_title)
    case = {
        'title': title,
        'position': max((x.get('position', 0) for x in suite['test_cases']), default=-1) + 1,
        'type': 'other', 'status': status, 'severity': 'undefined',
        'priority': 'medium', 'behavior': 'undefined', 'layer': 'unknown',
        'automation_status': 'manual', 'is_flaky': False,
        'to_be_automated': False, 'muted_case': False,
        'description': description, 'preconditions': preconditions,
        'steps': steps(pairs),
    }
    suite['test_cases'].append(case)
    added.append((suite_title, case))


add('Редактор публикаций', 'Возобновление и публикация сохранённого черновика',
    'Проверка, что сохранённый черновик открывается с тем же содержимым и может быть опубликован один раз.',
    'Пользователь авторизован; подготовлены уникальный заголовок и текст для новой публикации.', [
        ('Открыть редактор, ввести заголовок и текст, выбрать место публикации и нажать «Сохранить в черновики».', 'Редактор закрывается; запись появляется в «Мои публикации» → «Черновики».'),
        ('Открыть сохранённый черновик.', 'Заголовок, текст и выбранное место публикации сохранены; доступно продолжение редактирования.'),
        ('При необходимости дополнить текст и нажать «Сейчас».', 'Публикация появляется в выбранном месте; в списке черновиков нет второй копии опубликованной записи.'),
    ], status='draft')

add('Редактор публикаций', 'Появление запланированной публикации в заданное время',
    'Продолжение кейса о планировании: проверяется не только список запланированных, но и итоговая публикация.',
    'Пользователь авторизован; подготовлен уникальный заголовок. Текущее время и часовой пояс тестового окружения известны; выбрано ближайшее допустимое время с запасом на ожидание.', [
        ('Создать запись через «В точное время» и выбрать допустимое будущее время.', 'Запись находится в «Мои публикации» → «Запланированные» и до наступления времени не видна как опубликованная.'),
        ('Дождаться выбранного времени и обновить раздел публикаций и место назначения.', 'Запись появляется в выбранном месте и больше не числится запланированной.'),
        ('Повторно обновить страницу.', 'Запись опубликована один раз; дубликата нет.'),
    ], status='draft')

add('Опросы', 'Завершение опроса после указанного срока',
    'Проверка результата ограничения срока опроса для участника после наступления времени окончания.',
    'Есть два тестовых пользователя. Первый может создать публикацию с опросом; второй может открыть её и ещё не голосовал. Время окончания устанавливается на ближайшее допустимое будущее время.', [
        ('Первым пользователем создать опрос с конечной датой и временем и опубликовать запись.', 'До окончания срока второй пользователь видит варианты ответа и может голосовать.'),
        ('Дождаться времени окончания и обновить запись под вторым пользователем.', 'Опрос показывает завершённое состояние; отправка нового голоса недоступна.'),
        ('Открыть результаты опроса в доступном для этой настройки режиме.', 'Отображаются итоги уже поданных голосов, их количество не увеличивается после окончания срока.'),
    ], status='draft')

add('Комментарии к посту', 'Отмена редактирования комментария сохраняет исходные вложения',
    'Дополнение к проверке отмены редактирования: исходный опубликованный комментарий не должен измениться.',
    'Пользователь авторизован; под доступной публикацией создан его комментарий с текстом и допустимым вложением.', [
        ('Открыть редактирование своего комментария и изменить текст или состав вложений, не сохраняя.', 'Несохранённые изменения видны только в форме редактирования.'),
        ('Отменить редактирование.', 'Форма закрывается; опубликованный текст и исходное вложение остаются прежними.'),
        ('Обновить публикацию.', 'Комментарий по-прежнему содержит исходные текст и вложение; несохранённые изменения не появились.'),
    ])


def recount(data):
    all_suites = list(suites(data))
    all_cases = [c for s in all_suites for c in s.get('test_cases', [])]
    data['summary']['suite_count'] = len(all_suites)
    data['summary']['test_case_count'] = len(all_cases)
    data['summary']['step_count'] = sum(len(c.get('steps', [])) for c in all_cases)


recount(updated)
updated['exported_at'] = datetime.now().astimezone().isoformat()
json_path = OUT / 'Все кейсы — точечная актуализация.json'
json_path.write_text(json.dumps(updated, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.7)
sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.75)
sec.right_margin = Inches(.75)
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(9.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(4)
doc.styles['Title'].font.size = Pt(20)
doc.styles['Title'].font.bold = True
title_ppr = doc.styles['Title']._element.get_or_add_pPr()
for b in title_ppr.findall(qn('w:pBdr')):
    title_ppr.remove(b)
doc.styles['Heading 1'].font.size = Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)

doc.add_paragraph('Актуализация тест кейсов соцсети', 'Title')
doc.add_paragraph('Полный текст изменённых и новых кейсов. Светло жёлтым отмечен добавленный или изменённый текст внутри существующего кейса; светло зелёная метка обозначает новый кейс.')
doc.add_paragraph(f'В полном JSON: {updated["summary"]["test_case_count"]} кейсов. Изменено существующих: {len(changed)}. Добавлено новых: {len(added)}.')


def highlight(run, fill):
    rpr = run._element.get_or_add_rPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    rpr.append(shd)


def add_diff(paragraph, before, after):
    before = before or ''
    after = after or ''
    old_tokens = re.findall(r'\S+|\s+', before)
    new_tokens = re.findall(r'\S+|\s+', after)
    matcher = difflib.SequenceMatcher(None, old_tokens, new_tokens, autojunk=False)
    for tag, _, _, j1, j2 in matcher.get_opcodes():
        if tag == 'delete':
            continue
        value = ''.join(new_tokens[j1:j2])
        if not value:
            continue
        run = paragraph.add_run(value)
        if tag != 'equal':
            highlight(run, 'FFF3D7')


def field(label, old_value, value, new=False):
    p = doc.add_paragraph()
    p.add_run(label + '. ').bold = True
    if new:
        p.add_run(value or '—')
    else:
        add_diff(p, old_value or '', value or '')


def show_case(suite_title, c, old=None):
    new = old is None
    h = doc.add_paragraph(style='Heading 1')
    h.add_run(('Новый кейс  ' if new else f'{c["id"]}  ') + c['title'] + ('  [черновик]' if c.get('status') == 'draft' else ''))
    h.paragraph_format.keep_with_next = True
    if new:
        p = doc.add_paragraph()
        run = p.add_run('НОВЫЙ КЕЙС')
        run.bold = True
        run.font.color.rgb = RGBColor(22, 100, 64)
        highlight(run, 'E8F6ED')
    else:
        field('Название', old.get('title'), c['title'])
    p = doc.add_paragraph()
    p.add_run('Набор. ').bold = True
    p.add_run(suite_title)
    field('Описание', old.get('description') if old else '', c.get('description'), new)
    field('Предусловия', old.get('preconditions') if old else '', c.get('preconditions'), new)
    old_steps = old.get('steps', []) if old else []
    for i, step in enumerate(c['steps']):
        previous = old_steps[i] if i < len(old_steps) else {}
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.18)
        p.add_run(f'{i+1}. ').bold = True
        add_diff(p, previous.get('action', ''), step.get('action', '')) if not new else p.add_run(step.get('action', ''))
        p.add_run(' Ожидается: ')
        add_diff(p, previous.get('expected_result', ''), step.get('expected_result', '')) if not new else p.add_run(step.get('expected_result', ''))


for case_id in changed:
    for s in suites(updated):
        for c in s.get('test_cases', []):
            if c.get('id') == case_id:
                show_case(s['title'], c, changed[case_id])

for suite_title, c in added:
    show_case(suite_title, c)

doc.add_paragraph('Требуют отдельной сверки', 'Heading 1')
doc.add_paragraph('7572: сейчас ожидается доступ неучастника к посту скрытой группы по прямой ссылке. Без проверки прав на стенде этот ожидаемый результат не менялся.')
doc.add_paragraph('10360: кейс уже в статусе черновика; поведение при повторном добавлении того же файла не подтверждено. Не подменял неопределённый результат догадкой.')
doc.add_paragraph('10377: кейс без шагов, а его описание содержит пометку о несогласованном ожидаемом результате. Он оставлен без изменений до согласования поведения вставки текста с изображением.')
doc.add_paragraph('10383: кейс помечен как устаревший в исходном экспорте; менять его как действующий кейс не стал.')

docx_path = OUT / 'Все кейсы — изменения и новые кейсы.docx'
doc.save(docx_path)
print(json_path)
print(docx_path)
print('changed', list(changed), 'added', len(added), 'total', updated['summary']['test_case_count'])
